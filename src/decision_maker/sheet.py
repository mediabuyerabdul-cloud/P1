"""Read research sheets (.xlsx/.csv) you provide, and write decisions back.

You collect the data (stats, demographics, views); these readers map your
columns to known fields. Two shapes are supported:
  - a competitor table (F1 Niche Validator)
  - a "new channels in the last 60 days" table (F2 Demand vs Saturation)
"""

from __future__ import annotations

import csv
import io
import re

import openpyxl

# --- F1: competitor table ---------------------------------------------------
CHANNEL_ALIASES = {
    "channel name": "channel", "channel": "channel",
    "subscribes": "subscribers", "subscribers": "subscribers", "subs": "subscribers",
    "avg. views per video": "avg_views", "avg views per video": "avg_views", "avg views": "avg_views",
    "monthly views": "monthly_views",
    "monthly income": "monthly_income",
    "rpm": "rpm",
    "total views": "total_views", "total videos": "total_videos",
    "avg. monthly uploads": "uploads_per_month", "avg monthly uploads": "uploads_per_month",
    "avg. video length": "video_length", "avg video length": "video_length",
    "monetization status": "monetization", "monetization": "monetization",
    "relevance level": "relevance_given", "competitor stages": "stage_given", "competitor stage": "stage_given",
    "top geographies": "geo", "top gender": "gender", "top age": "age",
    "channel creation country": "country", "categories": "category", "category": "category",
}
CHANNEL_NUMERIC = {"subscribers", "avg_views", "monthly_views", "monthly_income", "rpm",
                   "total_views", "total_videos", "uploads_per_month"}

# --- F2: new-channels-in-60-days table --------------------------------------
NEW_ALIASES = {
    "past 60 days new channels created": "url", "channel": "url", "url": "url",
    "channel url": "url", "channel link": "url",
    "avg views in this channel in latest section": "avg_views", "avg views": "avg_views",
    "avg. views": "avg_views", "average views": "avg_views",
    "most popular view in this channel": "most_popular", "most popular view": "most_popular",
    "most popular": "most_popular", "top video views": "most_popular",
}
NEW_NUMERIC = {"avg_views", "most_popular"}


def _num(v):
    if v is None:
        return None
    s = re.sub(r"[^\d.]", "", str(v).replace(" ", ""))
    try:
        return float(s) if s else None
    except ValueError:
        return None


def _table_from_bytes(data: bytes, filename: str):
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return [list(r) for r in csv.reader(io.StringIO(text))]
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def _parse(table, aliases, numeric, required):
    header_idx, colmap = None, {}
    for i, row in enumerate(table[:12]):
        mapped = [aliases.get(str(c or "").strip().lower()) for c in row]
        if sum(1 for m in mapped if m) >= 2 and any(m == required for m in mapped):
            for j, m in enumerate(mapped):
                if m and m not in colmap:
                    colmap[m] = j
            header_idx = i
            break
    if header_idx is None:
        return []
    records = []
    for row in table[header_idx + 1:]:
        if not any(c not in (None, "") for c in row):
            continue
        rec = {}
        for field, j in colmap.items():
            val = row[j] if j < len(row) else None
            rec[field] = _num(val) if field in numeric else ("" if val is None else str(val).strip())
        if rec.get(required):
            records.append(rec)
    return records


def read_channels(data: bytes, filename: str) -> list[dict]:
    return _parse(_table_from_bytes(data, filename), CHANNEL_ALIASES, CHANNEL_NUMERIC, "channel")


def read_new_channels(data: bytes, filename: str) -> list[dict]:
    return _parse(_table_from_bytes(data, filename), NEW_ALIASES, NEW_NUMERIC, "url")


def write_xlsx(rows: list[dict], recommendation: dict) -> bytes:
    wb = openpyxl.Workbook()
    dec = wb.active
    dec.title = "Decisions"
    cols = ["channel", "relevance", "stage", "tier", "avg_views", "subscribers",
            "monthly_income", "rpm", "monetization", "reason"]
    dec.append([c.replace("_", " ").title() for c in cols])
    for r in rows:
        dec.append([r.get(c, "") for c in cols])
    summ = wb.create_sheet("Best Decision")
    summ.append(["Verdict", recommendation.get("verdict", "")])
    summ.append(["Why", recommendation.get("headline", recommendation.get("reason", ""))])
    for label, value in recommendation.get("analysis", []):
        summ.append([label, value])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
